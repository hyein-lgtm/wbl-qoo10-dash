"""큐텐 운영 대시보드 — FastAPI 서버.

환경변수
  QOO10_KEY       셀러 인증키 (없으면 데모 데이터로 동작)
  DASH_PASSWORD   접속 비밀번호 (아이디는 아무거나). 비우면 잠금 없음
  SYNC_MINUTES    자동 동기화 주기(분), 기본 60
  DELAY_DAYS      결제 후 이 일수가 지나도 미발송이면 '배송지연', 기본 3
  DB_PATH         SQLite 경로, 기본 ./qoo10.db (Railway는 볼륨 경로 권장)
"""
import os, json, sqlite3, random, secrets, threading, time, traceback
from datetime import datetime, timedelta
from pathlib import Path
from fastapi import FastAPI, Depends, HTTPException, Query, UploadFile, File, Request
from pydantic import BaseModel
from fastapi.responses import FileResponse
from fastapi.security import HTTPBasic, HTTPBasicCredentials
from qoo10 import Qoo10

KEY = os.getenv("QOO10_KEY", "").strip()
if not KEY:  # PC 실행용: 같은 폴더나 상위 폴더의 qoo10_key.txt
    for f in (Path(__file__).with_name("qoo10_key.txt"), Path(__file__).parent.parent / "qoo10_key.txt"):
        if f.exists():
            KEY = f.read_text(encoding="utf-8").strip(); break
PASSWORD = os.getenv("DASH_PASSWORD", "")
SYNC_MIN = int(os.getenv("SYNC_MINUTES", "60"))
DELAY_DAYS = int(os.getenv("DELAY_DAYS", "3"))
DB = os.getenv("DB_PATH", str(Path(__file__).with_name("qoo10.db")))
DEMO = not KEY

app = FastAPI(title="큐텐 운영 대시보드")
security = HTTPBasic(auto_error=bool(PASSWORD))


def auth(cred: HTTPBasicCredentials = Depends(security)):
    if PASSWORD and not (cred and secrets.compare_digest(cred.password, PASSWORD)):
        raise HTTPException(401, headers={"WWW-Authenticate": "Basic"})


# ── DB ───────────────────────────────────────────────
def db():
    c = sqlite3.connect(DB)
    c.row_factory = sqlite3.Row
    return c


with db() as c:
    c.executescript("""
    CREATE TABLE IF NOT EXISTS products(item_code TEXT PRIMARY KEY, seller_code TEXT, title TEXT,
        price REAL, retail_price REAL, qty INTEGER, status TEXT, raw TEXT, updated TEXT);
    CREATE TABLE IF NOT EXISTS orders(order_no TEXT PRIMARY KEY, pack_no TEXT, item_code TEXT, seller_code TEXT,
        title TEXT, option TEXT, qty INTEGER, amount REAL, order_date TEXT, pay_date TEXT, ship_date TEXT,
        status TEXT, carrier TEXT, tracking TEXT, buyer TEXT, raw TEXT, updated TEXT);
    CREATE TABLE IF NOT EXISTS settings(k TEXT PRIMARY KEY, v TEXT);
    CREATE TABLE IF NOT EXISTS registrations(id INTEGER PRIMARY KEY AUTOINCREMENT, created TEXT, seller_code TEXT,
        title TEXT, price REAL, ok INTEGER, gd_no TEXT, message TEXT, params TEXT);
    CREATE TABLE IF NOT EXISTS sync_log(id INTEGER PRIMARY KEY AUTOINCREMENT, started TEXT, finished TEXT,
        ok INTEGER, products INTEGER, orders INTEGER, message TEXT);
    """)


def pick(d, *keys, default=""):
    for k in keys:
        if d.get(k) not in (None, ""):
            return d[k]
    return default


def num(v):
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        return 0.0


def dt(v):
    """큐텐 날짜 문자열 → 'YYYY-MM-DD HH:MM:SS'"""
    if not v:
        return ""
    s = str(v).replace("T", " ").replace("/", "-")[:19]
    for f in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%Y%m%d"):
        try:
            return datetime.strptime(s[:len(datetime.now().strftime(f))], f).strftime("%Y-%m-%d %H:%M:%S")
        except ValueError:
            continue
    return s


STATUS = {"1": "배송요청", "2": "배송준비", "3": "배송중", "4": "배송완료", "5": "구매확정",
          "Paid": "배송요청", "On request": "배송요청", "Checking order": "배송준비",
          "On delivery": "배송중", "Delivered": "배송완료"}


def norm_status(v):
    v = str(v or "")
    for k, s in STATUS.items():
        if k.lower() in v.lower():
            return s
    return v or "미확인"


# ── 동기화 ────────────────────────────────────────────
_lock = threading.Lock()


def sync():
    if not _lock.acquire(blocking=False):
        return {"ok": False, "message": "이미 동기화 중"}
    started = datetime.now().isoformat(timespec="seconds")
    np = no = 0
    try:
        if DEMO:
            np, no = seed_demo()
            msg = "데모 데이터 생성 (QOO10_KEY 미설정)"
        else:
            q = Qoo10(KEY)
            now = datetime.now().isoformat(timespec="seconds")
            prods = q.products()
            for p in prods:  # 목록 API에 상품명·가격이 없으면 상세 조회로 보충
                if not pick(p, "ItemTitle", "ItemName"):
                    try:
                        obj = q.detail(pick(p, "ItemCode", "GdNo", "ItemNo")).get("ResultObject")
                        obj = obj[0] if isinstance(obj, list) and obj else obj
                        if isinstance(obj, dict):
                            p.update({k: v for k, v in obj.items() if v not in (None, "") and not p.get(k)})
                    except Exception:
                        pass
            with db() as c:
                for p in prods:
                    c.execute("REPLACE INTO products VALUES(?,?,?,?,?,?,?,?,?)", (
                        str(pick(p, "ItemCode", "GdNo", "ItemNo")), pick(p, "SellerCode"),
                        pick(p, "ItemTitle", "ItemName"), num(pick(p, "ItemPrice", "SellPrice", default=0)),
                        num(pick(p, "RetailPrice", default=0)), int(num(pick(p, "ItemQty", "Qty", default=0))),
                        pick(p, "ItemStatus", default="S2"), json.dumps(p, ensure_ascii=False), now))
            end = datetime.now()
            start = end - timedelta(days=int(os.getenv("ORDER_DAYS", "90")))
            ords, cur = [], start
            while cur < end:  # 큐텐 조회기간 제한 대비 30일 단위
                nxt = min(cur + timedelta(days=30), end)
                ords += q.orders(cur.strftime("%Y%m%d"), nxt.strftime("%Y%m%d"))
                cur = nxt + timedelta(days=1)
            with db() as c:
                for o in ords:
                    c.execute("REPLACE INTO orders VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", (
                        str(pick(o, "orderNo", "OrderNo")), str(pick(o, "packNo", "PackNo")),
                        str(pick(o, "itemCode", "ItemCode", "ItemNo")), pick(o, "sellerItemCode", "SellerItemCode", "SellerCode"),
                        pick(o, "itemTitle", "ItemTitle"), pick(o, "option", "Option", "OptionInfo"),
                        int(num(pick(o, "orderQty", "OrderQty", default=1))),
                        num(pick(o, "total", "Total", "orderPrice", "OrderPrice", default=0)),
                        dt(pick(o, "orderDate", "OrderDate")), dt(pick(o, "PaymentDate", "paymentDate", "orderDate")),
                        dt(pick(o, "ShippingDate", "shippingDate")), norm_status(pick(o, "shippingStatus", "ShippingStatus", "ShippingStat")),
                        pick(o, "DeliveryCompany", "deliveryCompany"), pick(o, "TrackingNo", "trackingNo"),
                        pick(o, "buyer", "Buyer", "receiver"), json.dumps(o, ensure_ascii=False), now))
            np, no = len(prods), len(ords)
            msg = f"API 조합: {q.combo}"
        ok = 1
    except Exception as e:
        ok, msg = 0, f"{e}"[:1500]
        traceback.print_exc()
    finally:
        _lock.release()
    with db() as c:
        c.execute("INSERT INTO sync_log(started,finished,ok,products,orders,message) VALUES(?,?,?,?,?,?)",
                  (started, datetime.now().isoformat(timespec="seconds"), ok, np, no, msg))
    return {"ok": bool(ok), "products": np, "orders": no, "message": msg}


def seed_demo():
    random.seed(7)
    names = [("BESTコレクション 耳つぼジュエリー", 2325, 3000), ("小顔 耳つぼシール 24粒", 1650, 2200),
             ("スワロフスキー 耳つぼジュエリー ピュアホワイト", 2980, 3800), ("五行 お守りブレスレット", 3480, 4500),
             ("耳つぼジュエリー お試しセット", 990, 1500), ("ゴールド 耳つぼ 48粒", 2750, 3500),
             ("ローズクォーツ 耳つぼジュエリー", 2480, 3200), ("耳つぼマップ付き スターターキット", 3980, 4980)]
    now = datetime.now()
    with db() as c:
        c.execute("DELETE FROM products"); c.execute("DELETE FROM orders")
        for i, (t, p, r) in enumerate(names):
            c.execute("REPLACE INTO products VALUES(?,?,?,?,?,?,?,?,?)",
                      (f"10{i+1:07d}", f"WBL-{i+1:03d}", t, p, r, random.choice([3, 8, 25, 60, 140]), "S2", "{}", now.isoformat()))
        n = 0
        for d in range(90, -1, -1):
            day = now - timedelta(days=d)
            mega = 1 if d in range(30, 36) else 0
            for _ in range(random.randint(0, 2) + mega * random.randint(8, 16)):
                i = random.choices(range(8), weights=[8, 3, 3, 2, 2, 1, 1, 1])[0]
                t, p, _r = names[i]
                q = random.choice([1, 1, 1, 2])
                pay = day.replace(hour=random.randint(0, 23), minute=random.randint(0, 59))
                age = (now - pay).days
                if age <= 1: st = "배송요청"
                elif age <= 3: st = random.choice(["배송요청", "배송준비", "배송중"])
                elif age <= 7: st = random.choice(["배송요청", "배송중", "배송완료", "배송완료"])
                else: st = random.choices(["배송완료", "구매확정", "배송요청"], weights=[3, 6, 0.3])[0]
                ship = "" if st in ("배송요청", "배송준비") else (pay + timedelta(days=random.choice([1, 2, 2, 3, 5]))).strftime("%Y-%m-%d %H:%M:%S")
                n += 1
                c.execute("REPLACE INTO orders VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                          (f"9{n:08d}", f"P{n:07d}", f"10{i+1:07d}", f"WBL-{i+1:03d}", t, "", q, p * q,
                           pay.strftime("%Y-%m-%d %H:%M:%S"), pay.strftime("%Y-%m-%d %H:%M:%S"), ship, st,
                           "ヤマト運輸" if ship else "", f"{random.randint(10**11, 10**12)}" if ship else "",
                           f"購入者{n}", "{}", now.isoformat()))
    return len(names), n


def scheduler():
    time.sleep(3)
    while True:
        sync()
        time.sleep(SYNC_MIN * 60)


threading.Thread(target=scheduler, daemon=True).start()


# ── 집계 ─────────────────────────────────────────────
def since(days):
    return (datetime.now() - timedelta(days=days)).strftime("%Y-%m-%d %H:%M:%S") if days else "0000"


def delay_info(o):
    pay = datetime.strptime(o["pay_date"], "%Y-%m-%d %H:%M:%S") if o["pay_date"] else None
    if not pay:
        return None, False
    if o["ship_date"]:
        took = (datetime.strptime(o["ship_date"], "%Y-%m-%d %H:%M:%S") - pay).days
        return took, took > DELAY_DAYS
    waited = (datetime.now() - pay).days
    return waited, waited > DELAY_DAYS and o["status"] in ("배송요청", "배송준비")


@app.get("/api/summary", dependencies=[Depends(auth)])
def summary(days: int = 30):
    with db() as c:
        s = c.execute("SELECT COUNT(*) n, COALESCE(SUM(amount),0) rev, COUNT(DISTINCT buyer) buyers FROM orders WHERE pay_date>=?", (since(days),)).fetchone()
        prev = c.execute("SELECT COALESCE(SUM(amount),0) rev FROM orders WHERE pay_date>=? AND pay_date<?", (since(days * 2), since(days))).fetchone() if days else None
        prods = c.execute("SELECT COUNT(*) n, SUM(qty<10) low FROM products").fetchone()
        open_orders = [dict(r) for r in c.execute("SELECT * FROM orders WHERE status IN ('배송요청','배송준비')")]
        last = c.execute("SELECT * FROM sync_log ORDER BY id DESC LIMIT 1").fetchone()
    delayed = sum(1 for o in open_orders if delay_info(o)[1])
    return {"demo": DEMO, "delay_days": DELAY_DAYS, "revenue": s["rev"], "orders": s["n"],
            "aov": s["rev"] / s["n"] if s["n"] else 0, "buyers": s["buyers"],
            "revenue_prev": prev["rev"] if prev else None, "products": prods["n"], "low_stock": prods["low"] or 0,
            "open_orders": len(open_orders), "delayed": delayed, "last_sync": dict(last) if last else None}


@app.get("/api/daily", dependencies=[Depends(auth)])
def daily(days: int = 30):
    with db() as c:
        rows = c.execute("SELECT substr(pay_date,1,10) d, SUM(amount) rev, COUNT(*) n FROM orders WHERE pay_date>=? GROUP BY d ORDER BY d",
                         (since(days or 3650),)).fetchall()
    return [dict(r) for r in rows]


@app.get("/api/products", dependencies=[Depends(auth)])
def products(days: int = 30, q: str = "", flag: str = ""):
    with db() as c:
        rows = c.execute("""SELECT p.*, COALESCE(SUM(o.qty),0) sold, COALESCE(SUM(o.amount),0) rev, COUNT(o.order_no) n
            FROM products p LEFT JOIN orders o ON o.item_code=p.item_code AND o.pay_date>=?
            GROUP BY p.item_code ORDER BY rev DESC""", (since(days),)).fetchall()
    out = []
    for r in rows:
        d = dict(r); d.pop("raw", None)
        d["discount"] = (1 - d["price"] / d["retail_price"]) * 100 if d["retail_price"] else 0
        d["daily_sold"] = d["sold"] / (days or 90)
        d["days_left"] = d["qty"] / d["daily_sold"] if d["daily_sold"] else None
        if q and q.lower() not in f"{d['title']} {d['seller_code']} {d['item_code']}".lower():
            continue
        if flag == "low" and d["qty"] >= 10: continue
        if flag == "nosale" and d["n"] > 0: continue
        if flag == "top" and d["rev"] == 0: continue
        out.append(d)
    return out


@app.get("/api/orders", dependencies=[Depends(auth)])
def orders(days: int = 30, status: str = "", delayed: bool = False, q: str = ""):
    with db() as c:
        rows = [dict(r) for r in c.execute("SELECT * FROM orders WHERE pay_date>=? ORDER BY pay_date DESC", (since(days),))]
    out = []
    for o in rows:
        o.pop("raw", None)
        o["elapsed"], o["is_delayed"] = delay_info(o)
        if status and o["status"] != status: continue
        if delayed and not (o["is_delayed"] and not o["ship_date"]): continue
        if q and q.lower() not in f"{o['order_no']} {o['title']} {o['buyer']} {o['tracking']}".lower(): continue
        out.append(o)
    return out[:2000]


@app.post("/api/sync", dependencies=[Depends(auth)])
def sync_now():
    return sync()


@app.get("/api/sync-log", dependencies=[Depends(auth)])
def sync_log():
    with db() as c:
        return [dict(r) for r in c.execute("SELECT * FROM sync_log ORDER BY id DESC LIMIT 50")]


# ── 상품 등록 ───────────────────────────────────────────
IMG_DIR = Path(DB).parent / "images"
IMG_DIR.mkdir(exist_ok=True)
AI_KEY = os.getenv("ANTHROPIC_API_KEY", "").strip()
CONTACT = os.getenv("CONTACT_INFO", "cx@wisland.co.kr").strip()  # A/S 연락처 통일값
AI_MODEL = os.getenv("ANTHROPIC_MODEL", "claude-sonnet-4-5")
REQUIRED = ["SecondSubCat", "ItemTitle", "SellerCode", "ContactInfo", "StandardImage",
            "ItemDescription", "ItemPrice", "ItemQty", "ShippingNo"]


def client():
    if DEMO:
        raise HTTPException(400, "데모 모드에서는 큐텐 호출을 할 수 없어요 (QOO10_KEY 필요)")
    return Qoo10(KEY)


@app.get("/api/items/{code}/detail", dependencies=[Depends(auth)])
def item_detail(code: str):
    if DEMO:
        return {"demo": True, "ResultObject": {"ItemCode": code, "ItemTitle": "주얼패치 베스트", "SellPrice": "6999.0000", "RetailPrice": "0.0000",
                "ItemQty": "100", "ImageUrl": "https://picsum.photos/400", "SecondSubCatCd": "320001861", "ContactInfo": "info@wisland.co.kr",
                "ShippingNo": "813346", "ProductionPlaceType": "2", "ProductionPlace": "KR", "AvailableDateType": "0", "AvailableDateValue": "3"}}
    try:
        return client().detail(code)
    except Exception as e:
        raise HTTPException(502, str(e))


@app.get("/api/delivery-groups", dependencies=[Depends(auth)])
def delivery_groups():
    if DEMO:
        return [{"ShippingNo": "813346", "ShippingFee": 300, "ShippingType": "F", "transcName": "デモ"}]
    try:
        return client().delivery_groups()
    except Exception as e:
        raise HTTPException(502, str(e))


@app.post("/api/images", dependencies=[Depends(auth)])
async def upload_image(request: Request, file: UploadFile = File(...)):
    ext = (Path(file.filename or "").suffix or ".jpg").lower()
    if ext not in (".jpg", ".jpeg", ".png", ".gif", ".webp"):
        raise HTTPException(400, "jpg/png/gif/webp 이미지만 올릴 수 있어요")
    data = await file.read()
    if len(data) > 10 * 1024 * 1024:
        raise HTTPException(400, "10MB 이하 이미지만 올릴 수 있어요")
    name = f"{datetime.now():%Y%m%d%H%M%S}_{secrets.token_hex(4)}{ext}"
    (IMG_DIR / name).write_bytes(data)
    proto = request.headers.get("x-forwarded-proto", request.url.scheme)
    base = os.getenv("PUBLIC_BASE_URL") or f"{proto}://{request.headers.get('host', request.url.netloc)}"
    return {"url": f"{base}/img/{name}"}


@app.get("/img/{name}")  # 큐텐이 이미지를 가져갈 수 있게 인증 없이 공개
def public_image(name: str):
    f = IMG_DIR / Path(name).name
    if not f.exists():
        raise HTTPException(404)
    return FileResponse(f)


class DraftIn(BaseModel):
    name_ko: str
    features: str = ""
    spec: str = ""
    usage: str = ""
    keywords: str = ""
    tone: str = ""


DRAFT_SYSTEM = """あなたはQoo10 Japanの化粧品・ジュエリー売場に精通した日本人EC担当者です。
ブランド: wbL公式 (韓国発のウェルネス・ビューティーブランド、耳つぼジュエリーが主力。クワイエットラグジュアリーの上品なトーン)。
韓国語の商品情報から、Qoo10出品用の日本語テキストを作ります。
ルール:
- 薬機法・景品表示法に抵触する表現は禁止 (痩せる、治る、効く、医療効果、No.1など根拠のない最上級、before/after断定)。
  「〜をサポート」「気分転換に」「セルフケアのお供に」など控えめな表現にする。
- ItemTitle は全角50文字以内。検索されやすい語 (耳つぼジュエリー、耳つぼシール 等) を自然に含める。記号の乱用禁止。
- PromotionName は全角20文字以内の短いキャッチ。
- ItemDescription はQoo10商品詳細用のシンプルなHTML (h3, p, ul/li, table のみ。style属性は最小限。script禁止)。
  構成: 導入文 → 特徴(箇条書き) → 商品仕様(表) → 使い方 → ご注意。
- Keyword は検索キーワードをカンマ区切りで最大10個。
必ず次のJSONだけを出力: {"ItemTitle":"","PromotionName":"","ItemDescription":"","Keyword":""}"""


@app.post("/api/ai/draft", dependencies=[Depends(auth)])
def ai_draft(d: DraftIn):
    if not AI_KEY:
        raise HTTPException(400, "ANTHROPIC_API_KEY 환경변수가 없어서 AI 초안을 만들 수 없어요")
    user = f"商品名(韓国語): {d.name_ko}\n特徴: {d.features}\n仕様・構成: {d.spec}\n使い方: {d.usage}\n入れたい検索語: {d.keywords}\nトーン要望: {d.tone}"
    import requests as rq
    r = rq.post("https://api.anthropic.com/v1/messages", timeout=90,
                headers={"x-api-key": AI_KEY, "anthropic-version": "2023-06-01", "content-type": "application/json"},
                json={"model": AI_MODEL, "max_tokens": 3000, "system": DRAFT_SYSTEM,
                      "messages": [{"role": "user", "content": user}]})
    if r.status_code != 200:
        raise HTTPException(502, f"AI 호출 실패 {r.status_code}: {r.text[:300]}")
    text = "".join(b.get("text", "") for b in r.json().get("content", []))
    try:
        return json.loads(text[text.index("{"): text.rindex("}") + 1])
    except ValueError:
        raise HTTPException(502, "AI 응답을 해석하지 못했어요: " + text[:300])


class OptIn(BaseModel):
    html: str = ""
    item_code: str = ""


OPT_SYSTEM = """あなたはQoo10 Japanの出品担当者です。商品詳細ページ(HTMLテキストと画像)から、購入者が選ぶ選択肢(バリエーション)を抽出します。
このショップの詳細ページは次の構成が多い:
- 商品ごとのセクション見出し「01 シンプル シャイン」のように【番号 + 商品名】
- そのセクション内の「色」欄に、選択肢ラベル「A (シルバー)」「D (ゴールド)」のように【英字コード + (色)】
- 1商品に色が複数あれば英字コードも複数(例: 01 シンプルシャイン → A=シルバー, D=ゴールド)。色がない商品はコード1つ。
やること:
- 英字コードごとに1件ずつ、それが属する商品名(見出し)と番号、色を対応付ける。
- 商品名・色はページ表記どおり(カタカナを変えない)。商品名内のスペースは詰めてよい。
- 選択肢でないもの(特徴、使い方、注意事項)は含めない。英字コードが無いページでは code を空にして商品・色ごとに1件。
- オプション名はページ表記を優先(なければ「デザイン」)。
必ず次のJSONだけを出力:
{"name":"デザイン","items":[{"code":"A","no":"01","product":"シンプルシャイン","color":"シルバー"}],"found_in":"image|text|none","note":""}"""


def _strip_html(h):
    import re, html as _h
    t = re.sub(r"(?is)<(script|style).*?</\1>", " ", h)
    t = re.sub(r"(?s)<[^>]+>", " ", t)
    return re.sub(r"\s+", " ", _h.unescape(t)).strip()


def _image_tiles(urls, width=1000, tile_h=1400, overlap=120):
    """이미지를 받아 가로 1000px로 줄이고, 세로로 긴 것은 1400px씩 잘라 base64 JPEG 목록으로."""
    import base64, io, requests as rq
    from PIL import Image
    out = []
    for u in urls:
        try:
            raw = rq.get(u, timeout=20, headers={"User-Agent": "Mozilla/5.0", "Referer": "https://www.qoo10.jp/"}).content
            im = Image.open(io.BytesIO(raw))
            im.seek(0)
            im = im.convert("RGB")
        except Exception:
            continue
        if im.width > width:
            im = im.resize((width, max(1, round(im.height * width / im.width))))
        y = 0
        while y < im.height:
            part = im.crop((0, y, im.width, min(im.height, y + tile_h)))
            if part.height > 40:
                buf = io.BytesIO()
                part.save(buf, "JPEG", quality=80)
                out.append(base64.b64encode(buf.getvalue()).decode())
            if y + tile_h >= im.height:
                break
            y += tile_h - overlap
    return out


@app.post("/api/ai/options", dependencies=[Depends(auth)])
def ai_options(d: OptIn):
    import re, requests as rq
    if not AI_KEY:
        raise HTTPException(400, "ANTHROPIC_API_KEY 환경변수가 없어서 옵션을 자동으로 읽을 수 없어요")
    html = d.html
    if not html.strip() and d.item_code:  # 기존 상품: 큐텐 상세 페이지에서 직접 가져오기
        try:
            r = rq.get(f"https://www.qoo10.jp/gmkt.inc/Goods/GoodsDetailInfo.aspx?goodscode={d.item_code}", timeout=30,
                       headers={"User-Agent": "Mozilla/5.0"})
            html = r.text
        except Exception as e:
            raise HTTPException(502, f"큐텐 상세 페이지를 가져오지 못했어요: {e}")
    if not html.strip():
        raise HTTPException(400, "상세 HTML이 비어 있어요")
    imgs = []
    for u in re.findall(r"""<img[^>]+src=["']([^"']+)["']""", html, flags=re.I):
        u = "https:" + u if u.startswith("//") else u
        if u.startswith("http") and u not in imgs and not re.search(r"(icon|blank|spacer|logo)", u, re.I):
            imgs.append(u)
    text = _strip_html(html)[:6000]
    best, notes = None, []
    fmt = os.getenv("OPTION_FORMAT", "{no}.{product} {code}({color})")
    tiles = _image_tiles(imgs[:80])  # 긴 상세 이미지를 잘라서 AI 크기 제한(2000px)을 피함
    if not tiles and not text:
        raise HTTPException(400, "읽을 수 있는 상세 이미지가 없어요")
    # 요청 크기 제한 대비: 이미지 조각 70개 또는 약 18MB씩 나눠서 순서대로 읽음 (앞 묶음 끝 3장은 겹쳐서 연결 유지)
    chunks, cur, size = [], [], 0
    for tl in tiles:
        if cur and (len(cur) >= 70 or size + len(tl) > 18_000_000):
            chunks.append(cur); cur = cur[-3:]; size = sum(map(len, cur))
        cur.append(tl); size += len(tl)
    if cur or not chunks:
        chunks.append(cur)
    items, name, notes = [], "", []
    for n, ch in enumerate(chunks, 1):
        content = [{"type": "text", "text": f"詳細ページのテキスト:\n{text or '(なし)'}\n\n以下は詳細ページの画像です(上から順番通り、長い画像は分割済み。{n}/{len(chunks)}部分)。見出しと色ラベルが別の画像に分かれていることがあります。"}]
        content += [{"type": "image", "source": {"type": "base64", "media_type": "image/jpeg", "data": b}} for b in ch]
        r = rq.post("https://api.anthropic.com/v1/messages", timeout=300,
                    headers={"x-api-key": AI_KEY, "anthropic-version": "2023-06-01", "content-type": "application/json"},
                    json={"model": AI_MODEL, "max_tokens": 4000, "system": OPT_SYSTEM,
                          "messages": [{"role": "user", "content": content}]})
        if r.status_code != 200:
            notes.append(f"{n}번째 묶음 실패 {r.status_code}: {r.text[:200]}")
            continue
        t = "".join(b.get("text", "") for b in r.json().get("content", []))
        try:
            jj = json.loads(t[t.index("{"): t.rindex("}") + 1])
        except ValueError:
            notes.append(f"{n}번째 묶음 응답 해석 실패")
            continue
        name = name or jj.get("name", "")
        items += jj.get("items") or [{"code": "", "product": v, "color": ""} for v in jj.get("values", [])]
    if not items and notes:
        raise HTTPException(502, " / ".join(notes))
    j = {"name": name, "items": items, "found_in": "image", "note": " / ".join(notes)}
    items = j.get("items") or [{"code": "", "product": v, "color": ""} for v in j.get("values", [])]
    items.sort(key=lambda it: 0 if it.get("product") else 1)  # 같은 코드면 상품명 있는 쪽 우선
    seen, rows = set(), []
    for it in items:
        code = str(it.get("code") or "").strip().upper()
        key = code or (it.get("product", ""), it.get("color", ""))
        if key in seen:
            continue
        seen.add(key)
        prod = str(it.get("product") or "").replace(" ", "").replace("\u3000", "")
        color = str(it.get("color") or "").strip()
        no = str(it.get("no") or "").strip()
        no_i = int("".join(ch for ch in no if ch.isdigit()) or 0) or 9999
        no_s = f"{no_i:02d}" if no_i != 9999 else ""
        v = fmt.format(code=code, no=no_s, product=prod, color=color)
        v = v.replace("()", "").replace("  ", " ").strip(" .")
        if not no_s:
            v = v.lstrip(".").strip()
        rows.append({"sort": (no_i, code or "~"), "value": v, **it})
    rows.sort(key=lambda x: x["sort"])
    return {"name": j.get("name") or "デザイン", "values": [x["value"] for x in rows],
            "items": [{k: v for k, v in x.items() if k != "sort"} for x in rows],
            "found_in": j.get("found_in", "image"), "note": j.get("note", ""), "images": len(imgs)}


class PriceMapIn(BaseModel):
    values: list
    url: str = "https://wblbeauty.com/product/list.html?cate_no=1105"
    base_krw: float = 17900
    base_jpy: float = 2400


PRICE_SYSTEM = """あなたは日韓EC担当者です。Qoo10のオプション(日本語)が、韓国自社モールのどの商品に当たるかを対応付けます。
入力: オプション一覧(番号付き)と、韓国モールの商品一覧テキスト(商品名・販売価格)。
- 名前の意味・モチーフ(例: スワール=스왈, パール=진주, ハート=하트, リボン=리본, 蝶=나비, クローバー=클로버)で最も近い商品を選ぶ。
- 販売価格は割引後の「판매가」(小さい方の価格)を使う。消費者価格は使わない。
- 自信がない場合は confidence を low に。該当なしは kr_name を空、kr_price を0に。
必ず次のJSONだけを出力: {"rows":[{"i":1,"kr_name":"","kr_price":0,"confidence":"high|mid|low"}]}"""


@app.post("/api/ai/price-map", dependencies=[Depends(auth)])
def ai_price_map(d: PriceMapIn):
    import requests as rq
    if not AI_KEY:
        raise HTTPException(400, "ANTHROPIC_API_KEY가 없어요")
    vals = [str(v).strip() for v in d.values if str(v).strip()]
    if not vals:
        raise HTTPException(400, "옵션 값이 없어요")
    try:
        page = rq.get(d.url, timeout=30, headers={"User-Agent": "Mozilla/5.0"}).text
    except Exception as e:
        raise HTTPException(502, f"자사몰 페이지를 가져오지 못했어요: {e}")
    text = _strip_html(page)[:15000]
    user = "オプション一覧:\n" + "\n".join(f"{i}. {v}" for i, v in enumerate(vals, 1)) + f"\n\n韓国モール商品一覧テキスト:\n{text}"
    r = rq.post("https://api.anthropic.com/v1/messages", timeout=120,
                headers={"x-api-key": AI_KEY, "anthropic-version": "2023-06-01", "content-type": "application/json"},
                json={"model": AI_MODEL, "max_tokens": 4000, "system": PRICE_SYSTEM, "messages": [{"role": "user", "content": user}]})
    if r.status_code != 200:
        raise HTTPException(502, f"AI 호출 실패 {r.status_code}: {r.text[:200]}")
    t = "".join(b.get("text", "") for b in r.json().get("content", []))
    try:
        rows = json.loads(t[t.index("{"): t.rindex("}") + 1])["rows"]
    except (ValueError, KeyError):
        raise HTTPException(502, "AI 응답 해석 실패: " + t[:200])
    ratio = d.base_jpy / d.base_krw
    out = []
    for row in rows:
        i = int(row.get("i", 0)) - 1
        if not 0 <= i < len(vals):
            continue
        kp = num(row.get("kr_price") or 0)
        add = round((kp * ratio - d.base_jpy) / 10) * 10 if kp else None
        out.append({"value": vals[i], "kr_name": row.get("kr_name", ""), "kr_price": kp, "add": add,
                    "confidence": row.get("confidence", "low"), "over_cap": add is not None and abs(add) > d.base_jpy * 0.5})
    return {"ratio": ratio, "rows": out}


# ── 폴더 일괄 등록 ─────────────────────────────────────
FX = float(os.getenv("FX_KRW_PER_JPY", "9.3"))  # 정가 표시용 환율 (100엔=930원)
_MALL_CACHE = {}


def _ai(system, user_content, max_tokens=2000):
    import requests as rq
    if not AI_KEY:
        raise HTTPException(400, "ANTHROPIC_API_KEY가 없어요")
    r = rq.post("https://api.anthropic.com/v1/messages", timeout=180,
                headers={"x-api-key": AI_KEY, "anthropic-version": "2023-06-01", "content-type": "application/json"},
                json={"model": AI_MODEL, "max_tokens": max_tokens, "system": system,
                      "messages": [{"role": "user", "content": user_content}]})
    if r.status_code != 200:
        raise HTTPException(502, f"AI 호출 실패 {r.status_code}: {r.text[:200]}")
    t = "".join(b.get("text", "") for b in r.json().get("content", []))
    try:
        return json.loads(t[t.index("{"): t.rindex("}") + 1])
    except ValueError:
        raise HTTPException(502, "AI 응답 해석 실패: " + t[:200])


def _mall_text(url):
    """자사몰 목록 페이지 → 상품번호 표시가 남은 텍스트 (10분 캐시)"""
    import re, requests as rq
    hit = _MALL_CACHE.get(url)
    if hit and time.time() - hit[0] < 600:
        return hit[1]
    page = rq.get(url, timeout=30, headers={"User-Agent": "Mozilla/5.0"}).text
    page = re.sub(r"""<a[^>]+product_no=(\d+)[^>]*>""", r" [#\1] ", page)
    text = _strip_html(page)[:20000]
    _MALL_CACHE[url] = (time.time(), text)
    return text


class BatchPriceIn(BaseModel):
    folder: str                      # 폴더 이름 (상품번호·한국어 이름 포함 가능)
    values: list = []                # 옵션 값 (없으면 단품)
    url: str = "https://wblbeauty.com/product/list.html?cate_no=1105"
    base_krw: float = 17900
    base_jpy: float = 2400


BATCH_PRICE_SYSTEM = """あなたは日韓EC担当者です。韓国自社モールの商品一覧テキスト([#番号]が商品番号)から価格を探します。
対象: (1) フォルダ名が指す商品 (2) 各オプション(日本語)が指す商品。
- フォルダ名に商品番号(例 7663)があればその番号の商品。なければ名前で最も近い商品。
- オプションは名前の意味・モチーフで対応(スワール=스왈, パール=진주, ハート=하트, リボン=리본, 蝶=나비, クローバー=클로버, フラワー=플라워 等)。
- price は割引後の「판매가」、consumer は「소비자가」(取り消し線の高い方)。
必ず次のJSONだけを出力:
{"main":{"no":"","kr_name":"","price":0,"consumer":0},"options":[{"i":1,"kr_name":"","price":0,"confidence":"high|mid|low"}]}"""


@app.post("/api/batch/price", dependencies=[Depends(auth)])
def batch_price(d: BatchPriceIn):
    try:
        text = _mall_text(d.url)
    except Exception as e:
        raise HTTPException(502, f"자사몰 페이지를 가져오지 못했어요: {e}")
    vals = [str(v) for v in d.values if str(v).strip()]
    user = f"フォルダ名: {d.folder}\nオプション:\n" + ("\n".join(f"{i}. {v}" for i, v in enumerate(vals, 1)) or "(なし)") + f"\n\n商品一覧:\n{text}"
    j = _ai(BATCH_PRICE_SYSTEM, user, 3000)
    ratio = d.base_jpy / d.base_krw
    to_jpy = lambda krw: int(round(krw * ratio / 10) * 10)
    main = j.get("main") or {}
    opts = []
    for o in j.get("options", []):
        i = int(o.get("i", 0)) - 1
        if 0 <= i < len(vals):
            kp = num(o.get("price") or 0)
            opts.append({"value": vals[i], "kr_name": o.get("kr_name", ""), "kr_price": kp,
                         "jpy": to_jpy(kp) if kp else None, "confidence": o.get("confidence", "low")})
    priced = [o["jpy"] for o in opts if o["jpy"]]
    item_price = min(priced) if priced else (to_jpy(num(main.get("price") or 0)) if main.get("price") else None)
    for o in opts:  # 판매가 = 최저 옵션가, 나머지는 추가금 (±50% 상한)
        if o["jpy"] and item_price:
            add = o["jpy"] - item_price
            cap = int(item_price * 0.5 // 10 * 10)
            o["add"], o["over_cap"] = min(add, cap), add > cap
        else:
            o["add"], o["over_cap"] = 0, False
    consumer = num(main.get("consumer") or 0)
    retail = int(consumer / FX // 100 * 100) if consumer else ""  # 정가: 소비자가 환율 환산 이하(경표법)
    if retail and item_price and retail <= item_price:
        retail = ""
    return {"main": main, "item_price": item_price, "retail_price": retail, "options": opts, "ratio": ratio}


class TitleIn(BaseModel):
    name_ko: str
    options: list = []
    keywords: str = ""


TITLE_SYSTEM = """あなたはQoo10 Japanの出品担当です。ショップ「wbL公式」の商品名を作ります。
形式: 「wbL 純正品 耳つぼジュエリー {デザイン名(カタカナ、詳細ページ表記)} {モチーフ・素材の検索語1〜2個} 貼るだけ ピアス穴不要」
- 全角50文字以内。オプションが複数なら「選べる{N}種」を入れる(Nは実際の数)。
- 薬機法・景表法に触れる語(小顔、痩せる、効果、No.1、最安)は禁止。
- promotion は【】付き全角12文字以内、効能表現なし。
必ず次のJSONだけを出力: {"title":"","promotion":""}"""


@app.post("/api/ai/title", dependencies=[Depends(auth)])
def ai_title(d: TitleIn):
    opts = [str(o) for o in d.options if str(o).strip()]
    user = f"韓国語の商品名/フォルダ名: {d.name_ko}\nオプション({len(opts)}件): {' / '.join(opts[:30]) or 'なし'}\n入れたい語: {d.keywords}"
    return _ai(TITLE_SYSTEM, user, 500)


# ── 정리표(엑셀) 일괄 등록 ────────────────────────────────
# 자사몰 판매가·소비자가 (2026-09-30 wblbeauty.com cate_no=1105 기준). /api/mall/prices?refresh=1 로 갱신
MALL_DEFAULT = {
    "7677": (12900, 19900), "7031": (21900, 36900), "7236": (22900, 29900), "7201": (19900, 28900), "7778": (22900, 40900),
    "7780": (22900, 37900), "7779": (22900, 40900), "7188": (12900, 19900), "7257": (26900, 37900), "7439": (20900, 36900),
    "7235": (26900, 34900), "7191": (24900, 33900), "7341": (25900, 45900), "7187": (17900, 26400), "7662": (21900, 43900),
    "7663": (20900, 30900), "7313": (24900, 33900), "7777": (22900, 37900), "7297": (26900, 34900), "7221": (14900, 23900),
    "7395": (29900, 49900), "7225": (6900, 12900), "7661": (21900, 37900), "7189": (18400, 26900), "7247": (17900, 25900),
    "7242": (17900, 25900), "7314": (25900, 33900), "7658": (24900, 39900), "7776": (24900, 42900), "7305": (23900, 31900),
    "7344": (22900, 39900), "7659": (21900, 35900), "7660": (20900, 34900), "7375": (19900, 35900), "7433": (23900, 45900),
    "7343": (22900, 39900), "7190": (19900, 26900), "7345": (21900, 39900), "7306": (22900, 39900), "7296": (26900, 34900),
    "7307": (26900, 34900), "7200": (4900, 9900)}
_MALL_LIVE = {}


@app.get("/api/mall/prices", dependencies=[Depends(auth)])
def mall_prices(refresh: bool = False, url: str = "https://wblbeauty.com/product/list.html?cate_no=1105"):
    if refresh:
        text = _mall_text(url)
        j = _ai("韓国モールの商品一覧テキストから、[#番号]ごとに販売価格(割引後)と消費者価格を抜き出す。"
                "必ずJSONのみ: {\"items\":[{\"no\":\"7663\",\"price\":20900,\"consumer\":30900}]}", text, 4000)
        _MALL_LIVE.clear()
        _MALL_LIVE.update({str(i["no"]): (num(i.get("price")), num(i.get("consumer"))) for i in j.get("items", []) if i.get("no")})
    src = _MALL_LIVE or MALL_DEFAULT
    return {"source": "live" if _MALL_LIVE else "default(2026-09-30)", "fx": FX,
            "prices": {k: {"price": v[0], "consumer": v[1]} for k, v in src.items()}}


def _hdr(row):
    return [str(c or "").strip() for c in row]


def _col(h, *names):
    for n in names:
        for i, x in enumerate(h):
            if x == n or x.startswith(n):
                return i
    return None


@app.post("/api/sheet/parse", dependencies=[Depends(auth)])
async def sheet_parse(file: UploadFile = File(...)):
    import io, re
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(await file.read()), data_only=True, read_only=True)
    S = lambda n: [r for r in wb[n].iter_rows(values_only=True) if any(v not in (None, "") for v in r)] if n in wb.sheetnames else []
    v = lambda r, i: ("" if i is None or i >= len(r) or r[i] is None else str(r[i]).strip())
    split = lambda s: [x.strip() for x in re.split(r"\s*(?:\n|\$\$| / |／)\s*", s or "") if x.strip()]
    items = []
    # QSM 초안: 상품명 → 상세 HTML
    qsm = {}
    rows = S("QSM일괄등록_초안")
    if rows:
        h = _hdr(rows[0]); ci, cd, cm = _col(h, "E item_name"), _col(h, "V item_description"), _col(h, "AI item_material")
        for r in rows[1:]:
            qsm[v(r, ci)] = {"desc": v(r, cd), "material": v(r, cm)}
    # 옵션
    opts = {}
    rows = S("옵션")
    if rows:
        h = _hdr(rows[0]); c = {k: _col(h, k) for k in ["상품번호", "옵션명1", "값1", "옵션명2", "값2", "추가금"]}
        for r in rows[1:]:
            no = v(r, c["상품번호"])
            if not v(r, c["값1"]):
                continue
            o = opts.setdefault(no, {"name": v(r, c["옵션명1"]), "name2": v(r, c["옵션명2"]), "rows": []})
            o["rows"].append({"value": v(r, c["값1"]), "value2": v(r, c["값2"]), "add": num(v(r, c["추가금"]) or 0)})
    # 단품
    rows = S("단품")
    if rows:
        h = _hdr(rows[0])
        c = {k: _col(h, k) for k in ["상품번호", "한국어명", "큐텐 상품코드", "상태", "상품명(새)", "広告文", "検索ワード", "지금 판매가",
                                     "원산국", "素材", "대표이미지", "추가이미지", "상세 HTML", "카테고리 코드", "브랜드 코드"]}
        for r in rows[1:]:
            no, title = v(r, c["상품번호"]), v(r, c["상품명(새)"])
            if not no or not title:
                continue
            q = qsm.get(title, {})
            items.append({"key": f"s{no}", "kind": "단품", "no": no, "kr_name": v(r, c["한국어명"]), "qoo_code": v(r, c["큐텐 상품코드"]),
                          "status": v(r, c["상태"]), "title": title, "promo": v(r, c["広告文"]), "keywords": split(v(r, c["検索ワード"])),
                          "price_now": num(v(r, c["지금 판매가"]) or 0), "origin": v(r, c["원산국"]), "material": v(r, c["素材"]) or q.get("material", ""),
                          "main": v(r, c["대표이미지"]), "gallery": split(v(r, c["추가이미지"])), "desc": q.get("desc", ""),
                          "desc_path": v(r, c["상세 HTML"]).replace("전달 폴더", "").strip(), "category": v(r, c["카테고리 코드"]),
                          "brand": v(r, c["브랜드 코드"]), "opt": opts.get(no), "mode": "new"})
    # 모음전
    lines = {}
    rows = S("모음전")
    if rows:
        h = _hdr(rows[0]); c = {k: _col(h, k) for k in ["모음전", "글자", "상품번호", "옵션명", "선택지 문구", "추가금"]}
        for r in rows[1:]:
            nm = v(r, c["모음전"])
            if nm and v(r, c["선택지 문구"]):
                lines.setdefault(nm, {"name": v(r, c["옵션명"]) or "デザイン", "rows": []})["rows"].append(
                    {"value": v(r, c["선택지 문구"]), "letter": v(r, c["글자"]), "no": v(r, c["상품번호"])})
    rows = S("모음전_리스팅")
    if rows:
        h = _hdr(rows[0])
        c = {k: _col(h, k) for k in ["모음전", "큐텐 상품코드", "할 일", "상품명(새)", "広告文", "検索ワード", "대표이미지", "추가이미지", "상세 HTML", "원산국"]}
        base = next((i for i in items if i["category"]), {})
        for r in rows[1:]:
            nm = v(r, c["모음전"])
            if not nm:
                continue
            code = v(r, c["큐텐 상품코드"])
            items.append({"key": f"c{len(items)}", "kind": "모음전", "no": "", "kr_name": nm, "qoo_code": code if code.isdigit() else "",
                          "status": v(r, c["할 일"]), "title": v(r, c["상품명(새)"]), "promo": v(r, c["広告文"]),
                          "keywords": split(v(r, c["検索ワード"])), "price_now": 0, "origin": v(r, c["원산국"]),
                          "material": base.get("material", ""), "main": v(r, c["대표이미지"]), "gallery": split(v(r, c["추가이미지"])),
                          "desc": "", "desc_path": v(r, c["상세 HTML"]).replace("전달 폴더", "").strip(),
                          "category": base.get("category", ""), "brand": base.get("brand", ""), "coll": lines.get(nm),
                          "mode": "update" if code.isdigit() else "new"})
    dc = next((i["category"] for i in items if i["category"]), "")
    db_ = next((i["brand"] for i in items if i["brand"]), "")
    for i in items:
        i["category"] = i["category"] or dc
        i["brand"] = i["brand"] or db_
    return {"items": items, "sheets": wb.sheetnames}


# ── 간편 신규 등록 (자사몰 상품번호 + 폴더) ────────────────────
RULES_DEFAULT = {
    "title_template": "【公式】耳つぼジュエリー {line} 1箱20粒（1シート） {variant} 耳ツボ 耳つぼシール 貼るピアス 穴あけ不要 ギフト",
    "base_keywords": "耳ツボジュエリー, 耳ツボシール, 耳つぼピアス, ジュエリーシール, 耳つぼシート, イヤージュエリー, 耳ツボピアス, シールピアス, プレゼント",
    "banned_words": "小顔, 痩せる, ダイエット効果, 治る, 効く, 効果, No.1, 最安, 医療",
    "category": "320001861", "brand": "139162", "shipping": "813346", "contact": "cx@wisland.co.kr",
    "base_krw": 17900, "base_jpy": 2400, "fx": 9.3, "opt_qty": 10, "single_qty": 30, "prefix": "WBL-",
    "mall_detail_url": "https://wblbeauty.com/product/detail.html?product_no={no}",
    "max_title": 100, "detail_mb_warn": 40,
}


def rules():
    with db() as c:
        row = c.execute("SELECT v FROM settings WHERE k='rules'").fetchone()
    r = dict(RULES_DEFAULT)
    if row:
        try:
            r.update(json.loads(row["v"]))
        except ValueError:
            pass
    return r


@app.get("/api/settings", dependencies=[Depends(auth)])
def get_settings():
    return rules()


@app.post("/api/settings", dependencies=[Depends(auth)])
def save_settings(body: dict):
    r = {k: body[k] for k in RULES_DEFAULT if k in body}
    cur = rules(); cur.update(r)
    with db() as c:
        c.execute("REPLACE INTO settings VALUES('rules',?)", (json.dumps(cur, ensure_ascii=False),))
    return cur


def _mall_detail(no, url_tpl):
    """자사몰 상세 페이지 → (텍스트, 옵션 select 요약, 메타 가격)"""
    import re, requests as rq
    page = rq.get(url_tpl.format(no=no), timeout=30, headers={"User-Agent": "Mozilla/5.0"}).text
    metas = dict(re.findall(r'<meta[^>]+property="(product:[^"]+|og:title)"[^>]+content="([^"]*)"', page))
    selects = []
    for attrs, body in re.findall(r"<select([^>]*)>(.*?)</select>", page, flags=re.S | re.I):
        t = re.search(r'option_title="([^"]+)"', attrs)
        opts = [_strip_html(x) for x in re.findall(r"<option[^>]*>(.*?)</option>", body, flags=re.S)]
        opts = [o for o in opts if o and not re.search(r"선택|-{3,}|필수", o)]
        if t and opts:
            selects.append({"title": t.group(1), "values": opts})
    return _strip_html(page)[:6000], selects, metas


DRAFT_NEW_SYSTEM = """あなたはQoo10 Japan「wbL公式」の出品担当です。韓国自社モールの商品情報と、日本語の詳細ページテキストから出品データを作ります。
- line: 日本語のライン名(カタカナ)。詳細ページに表記があればそれを最優先で同じ表記に。例: セリン パール クローバー
- options: 韓国モールの選択肢を日本語に。オプション名は カラー/サイズ/タイプ のいずれか。値は詳細ページの表記を優先(例: シャーベットホワイト, S_3mm)。
  色+サイズのように2段なら2つ。選択肢がなければ空配列。「1箱20粒」等の数量は選択肢ではない。
- promo: 広告文、全角20文字以内、効能表現なし(例: パールとクリスタル 穴あけ不要)
- motif: 検索用のモチーフ語0〜2個(例: 花, 蝶, 星, 月, バラ, キラキラ, パール)
- kr_name, price(판매가), consumer(소비자가) を韓国モールから。
必ずJSONのみ: {"kr_name":"","price":0,"consumer":0,"line":"","options":[{"name":"カラー","values":[""]}],"promo":"","motif":[]}"""


class NewDraftIn(BaseModel):
    no: str
    html: str = ""


@app.post("/api/new/draft", dependencies=[Depends(auth)])
def new_draft(d: NewDraftIn):
    import re
    R = rules()
    no = re.sub(r"\D", "", d.no)
    if not no:
        raise HTTPException(400, "자사몰 상품번호가 필요해요")
    warn = []
    try:
        text, selects, metas = _mall_detail(no, R["mall_detail_url"])
    except Exception as e:
        text, selects, metas = "", [], {}
        warn.append(f"자사몰 페이지를 못 읽었어요({e.__class__.__name__}) — 저장된 가격표로 대신해요")
    jp_text = _strip_html(d.html)[:4000] if d.html else ""
    user = (f"商品番号: {no}\n韓国モールのメタ: {json.dumps(metas, ensure_ascii=False)}\n韓国モールの選択肢: {json.dumps(selects, ensure_ascii=False)}\n"
            f"韓国モール本文: {text[:4000]}\n\n日本語の詳細ページ本文: {jp_text or '(なし)'}")
    j = _ai(DRAFT_NEW_SYSTEM, user, 2000)
    price_kr, cons_kr = num(j.get("price") or 0), num(j.get("consumer") or 0)
    if not price_kr and no in MALL_DEFAULT:
        price_kr, cons_kr = MALL_DEFAULT[no]
    if not price_kr:
        warn.append("자사몰 판매가를 찾지 못했어요 — 판매가를 직접 입력해 주세요")
    ratio = float(R["base_jpy"]) / float(R["base_krw"])
    price = int(round(price_kr * ratio / 10) * 10) if price_kr else 0
    retail = int(cons_kr / float(R["fx"]) // 100 * 100) if cons_kr else 0
    if retail and retail <= price:
        retail = 0
    opts = [o for o in (j.get("options") or []) if o.get("values")][:2]
    # 상품명 규칙
    def variant():
        parts = []
        for o in opts:
            n, vs = o.get("name", ""), o["values"]
            if "カラー" in n:
                parts.append(" ".join(vs) if len(vs) <= 2 and all(len(x) <= 5 for x in vs) else f"全{len(vs)}色")
            elif "サイズ" in n:
                parts.append(f"{len(vs)}サイズ")
            elif len(vs) > 1:
                parts.append(f"{len(vs)}タイプ")
        return " ".join(parts)
    line = (j.get("line") or "").strip()
    title = re.sub(r"\s+", " ", R["title_template"].format(line=line, variant=variant())).strip()
    if len(title) > int(R["max_title"]):
        warn.append(f"상품명이 {len(title)}자로 길어요 (최대 {R['max_title']})")
    keywords = [k.strip() for k in (j.get("motif") or []) + R["base_keywords"].split(",") if k.strip()]
    promo = (j.get("promo") or "").strip()
    banned = [w.strip() for w in R["banned_words"].split(",") if w.strip()]
    hit = [w for w in banned if w in title + promo]
    if hit:
        warn.append("금지어 포함: " + ", ".join(hit))
    rows = []
    if len(opts) == 2:
        for v1 in opts[0]["values"]:
            for v2 in opts[1]["values"]:
                rows.append({"value": v1, "value2": v2, "price": 0})
    elif opts:
        rows = [{"value": v, "price": 0} for v in opts[0]["values"]]
    return {"no": no, "kr_name": j.get("kr_name", ""), "line": line, "title": title, "promo": promo, "keywords": keywords,
            "price": price, "retail": retail, "kr_price": price_kr, "kr_consumer": cons_kr,
            "opt": {"name": opts[0]["name"] if opts else "", "name2": opts[1]["name"] if len(opts) > 1 else "", "rows": rows},
            "warn": warn, "rules": {k: R[k] for k in ("category", "brand", "shipping", "contact", "opt_qty", "single_qty", "prefix", "detail_mb_warn")}}


class RegisterIn(BaseModel):
    params: dict
    confirm: bool = False
    gallery: list = []   # 추가이미지(갤러리) URL, 최대 50
    options: dict = {}   # {"name": "デザイン", "rows": [{"value","price","qty","code"}]}


def apply_gallery(item_code, seller_code, urls):
    urls = [u for u in urls if u][:50]
    if not urls:
        return None
    p = {"ItemCode": str(item_code), "SellerCode": seller_code or ""}
    for i, u in enumerate(urls, 1):
        p[f"EnlargedImage{i}"] = u
    if DEMO:
        return {"ResultCode": 0, "ResultMsg": "DEMO", "count": len(urls)}
    return client().write("ItemsContents.EditGoodsMultiImage", p)


def apply_options(item_code, seller_code, opt):
    """rows: [{value, price, qty, code}] (1단) 또는 [{value, value2, ...}] + opt.name2 (2단)"""
    rows = [r for r in (opt or {}).get("rows", []) if str(r.get("value", "")).strip()]
    name = str((opt or {}).get("name") or "オプション").strip()
    name2 = str((opt or {}).get("name2") or "").strip()
    if not rows:
        return None
    two = bool(name2) and any(str(r.get("value2", "")).strip() for r in rows)
    def fmt(r, style):
        v, pr, q, cd = r["value"], str(r.get("price") or 0), str(r.get("qty") or 0), r.get("code") or ""
        if two:
            v2 = r.get("value2", "")
            return f"{name}||*{v}||*{name2}||*{v2}||*{pr}||*{q}||*{cd}"
        return (f"{name}||*{v}||*{pr}||*{q}||*{cd}" if style == 1
                else f"{name}||*{v}||*||*||*{pr}||*{q}||*{cd}")
    if DEMO:
        return {"ResultCode": 0, "ResultMsg": "DEMO", "format": 1, "InventoryInfo": "$$".join(fmt(r, 1) for r in rows)[:300]}
    tried = []
    for style in ((1,) if two else (1, 2)):  # 큐텐 문서의 표기를 순서대로 시도
        info = "$$".join(fmt(r, style) for r in rows)
        res = client().write("ItemsOptions.EditGoodsInventory",
                             {"ItemCode": str(item_code), "SellerCode": seller_code or "", "InventoryInfo": info})
        res["format"] = style
        tried.append(res)
        if str(res.get("ResultCode")) == "0":
            return res
    return {"ResultCode": -1, "ResultMsg": "옵션 등록 실패 (형식 거절)", "tried": tried}


def _ok(res):
    return res is None or str(res.get("ResultCode")) == "0"


class PatchIn(BaseModel):
    seller_code: str = ""
    gallery: list = []
    options: dict = {}


@app.post("/api/items/{code}/patch", dependencies=[Depends(auth)])
def patch_item(code: str, body: PatchIn):
    """이미 등록된 상품에 추가이미지·옵션을 넣는다."""
    try:
        g = apply_gallery(code, body.seller_code, body.gallery)
        o = apply_options(code, body.seller_code, body.options)
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(502, str(e))
    return {"ok": _ok(g) and _ok(o), "gallery": g, "options": o}


@app.post("/api/items/register", dependencies=[Depends(auth)])
def register_item(body: RegisterIn):
    p = {k: str(v).strip() for k, v in body.params.items() if str(v).strip() != ""}
    p.setdefault("ContactInfo", CONTACT)
    missing = [k for k in REQUIRED if k not in p]
    if missing:
        raise HTTPException(400, "필수 항목이 비어 있어요: " + ", ".join(missing))
    if not body.confirm:
        raise HTTPException(400, "최종 확인 체크가 필요해요")
    with db() as c:  # 같은 셀러코드 중복 등록 방지
        dup = c.execute("SELECT gd_no FROM registrations WHERE seller_code=? AND ok=1", (p["SellerCode"],)).fetchone()
    if dup:
        raise HTTPException(409, f"셀러코드 {p['SellerCode']}는 이미 등록됐어요 (상품번호 {dup['gd_no']}). 셀러코드를 바꿔 주세요")
    if DEMO:
        res = {"ResultCode": 0, "ResultMsg": "DEMO", "ResultObject": {"GdNo": f"DEMO{secrets.randbelow(10**6):06d}"}}
    else:
        res = client().write("ItemsBasic.SetNewGoods", p)
    ok = str(res.get("ResultCode")) == "0"
    obj = res.get("ResultObject")
    gd = str((obj or {}).get("GdNo", "")) if isinstance(obj, dict) else str(obj or "")
    msg = res.get("ResultMsg") or res.get("ErrorMsg") or ""
    with db() as c:
        c.execute("INSERT INTO registrations(created,seller_code,title,price,ok,gd_no,message,params) VALUES(?,?,?,?,?,?,?,?)",
                  (datetime.now().isoformat(timespec="seconds"), p["SellerCode"], p["ItemTitle"], num(p["ItemPrice"]),
                   int(ok), gd, json.dumps(res, ensure_ascii=False)[:1500], json.dumps(p, ensure_ascii=False)))
    g = o = None
    if ok and gd:
        try:
            g = apply_gallery(gd, p["SellerCode"], body.gallery)
        except Exception as e:
            g = {"ResultCode": -1, "ResultMsg": str(e)}
        try:
            o = apply_options(gd, p["SellerCode"], body.options)
        except Exception as e:
            o = {"ResultCode": -1, "ResultMsg": str(e)}
    if ok and not DEMO:
        threading.Thread(target=sync, daemon=True).start()
    return {"ok": ok, "gd_no": gd, "message": msg, "raw": res, "gallery": g, "options": o}


class UpdateIn(BaseModel):
    seller_code: str = ""
    basic: dict | None = None        # 상품명·카테고리 등 기본정보 (전체 세트)
    price_qty: dict | None = None    # {"price","qty"}
    image: str | None = None         # 대표 이미지 URL
    description: str | None = None   # 상세 HTML
    gallery: list = []
    options: dict = {}


def _w(method, params):
    if DEMO:
        return {"ResultCode": 0, "ResultMsg": "DEMO", "method": method}
    try:
        return client().write(method, params)
    except HTTPException:
        raise
    except Exception as e:
        return {"ResultCode": -1, "ResultMsg": str(e)}


# UpdateGoods는 바꾸지 않는 필수 항목까지 전부 보내야 해서, 현재 상품 정보를 읽어 채운다
BASIC_MAP = {  # UpdateGoods 파라미터: 상세 조회 응답에서 찾을 키 후보
    "SecondSubCat": ["SecondSubCatCd", "SecondSubCat"], "ItemTitle": ["ItemTitle"], "PromotionName": ["PromotionName"],
    "ProductionPlaceType": ["ProductionPlaceType"], "ProductionPlace": ["ProductionPlace", "ProductionPlaceCd"],
    "AdultYN": ["AdultYN"], "ContactInfo": ["ContactInfo"], "ShippingNo": ["ShippingNo", "DeliveryGroupNo"],
    "Weight": ["Weight"], "AvailableDateType": ["AvailableDateType"], "AvailableDateValue": ["AvailableDateValue"],
    "Keyword": ["Keyword", "SearchKeyword"], "BrandNo": ["BrandNo"], "ModelNM": ["ModelNM", "ModelName"],
    "Material": ["Material"], "IndustrialCodeType": ["IndustrialCodeType"], "IndustrialCode": ["IndustrialCode"],
    "ManufactureDate": ["ManufactureDate"], "RetailPrice": ["RetailPrice"], "ExpireDate": ["ExpireDate"],
}
BASIC_DEFAULT = {"ProductionPlaceType": "2", "ProductionPlace": "KR", "AdultYN": "N", "AvailableDateType": "0",
                 "AvailableDateValue": "3", "ContactInfo": CONTACT}


def _clean(v):
    s = str(v).strip()
    if s.replace(".", "", 1).isdigit() and s.endswith(".0000"):
        s = s[:-5]
    return s


def update_basic(code, base, edits):
    import re
    cur = {}
    if not DEMO:
        try:
            obj = client().detail(code).get("ResultObject")
            obj = obj[0] if isinstance(obj, list) and obj else (obj or {})
            for k, cands in BASIC_MAP.items():
                for c in cands:
                    if obj.get(c) not in (None, ""):
                        cur[k] = _clean(obj[c]); break
        except Exception:
            pass
    if cur.get("RetailPrice") in ("0", "0.0"):
        cur.pop("RetailPrice")
    p = {**base, **{k: v for k, v in BASIC_DEFAULT.items() if v}, **cur,
         **{k: str(v).strip() for k, v in edits.items() if str(v).strip() != ""}}
    res = _w("ItemsBasic.UpdateGoods", p)
    for _ in range(4):  # 'XXXは必須です' 이면 기본값을 채워 재시도
        m = re.search(r"([A-Za-z]+)は必須", str(res.get("ResultMsg") or res.get("ErrorMsg") or ""))
        if not m or str(res.get("ResultCode")) == "0":
            break
        k = m.group(1)
        if k in p or k not in BASIC_DEFAULT or not BASIC_DEFAULT[k]:
            res["ResultMsg"] = f"{res.get('ResultMsg')} (필수 항목 {k} 값이 필요해요)"
            break
        p[k] = BASIC_DEFAULT[k]
        res = _w("ItemsBasic.UpdateGoods", p)
    res["sent"] = {k: (v if len(str(v)) < 80 else str(v)[:80] + "…") for k, v in p.items()}
    return res


@app.post("/api/items/{code}/update", dependencies=[Depends(auth)])
def update_item(code: str, body: UpdateIn):
    """등록된 상품 수정: 바뀐 부분만 해당 API로 보낸다."""
    base = {"ItemCode": str(code), "SellerCode": body.seller_code or ""}
    out = {}
    if body.basic:
        out["basic"] = update_basic(code, base, body.basic)
    if body.price_qty:
        out["price_qty"] = _w("ItemsOrder.SetGoodsPriceQty", {**base, "Price": str(body.price_qty.get("price", "")),
                                                               "Qty": str(body.price_qty.get("qty", ""))})
    if body.image:
        out["image"] = _w("ItemsContents.EditGoodsImage", {**base, "StandardImage": body.image})
    if body.description is not None:
        out["description"] = _w("ItemsContents.EditGoodsContents", {**base, "Contents": body.description})
    try:
        g = apply_gallery(code, body.seller_code, body.gallery)
        if g: out["gallery"] = g
        o = apply_options(code, body.seller_code, body.options)
        if o: out["options"] = o
    except Exception as e:
        out["gallery_options"] = {"ResultCode": -1, "ResultMsg": str(e)}
    if not out:
        raise HTTPException(400, "바뀐 내용이 없어요")
    if not DEMO:
        threading.Thread(target=sync, daemon=True).start()
    return {"ok": all(str(v.get("ResultCode")) == "0" for v in out.values()), "steps": out}


@app.post("/api/bulk/contact", dependencies=[Depends(auth)])
def bulk_contact():
    """판매중 전체 상품의 A/S 연락처를 CONTACT 값으로 통일"""
    with db() as c:
        codes = [r["item_code"] for r in c.execute("SELECT item_code FROM products")]
    if not codes:
        raise HTTPException(400, "상품 목록이 비어 있어요. 먼저 동기화해 주세요")
    results = []
    for code in codes:
        res = update_basic(code, {"ItemCode": code, "SellerCode": ""}, {"ContactInfo": CONTACT})
        results.append({"item_code": code, "ok": str(res.get("ResultCode")) == "0",
                        "message": res.get("ResultMsg") or res.get("ErrorMsg") or ""})
        time.sleep(0.3)
    return {"contact": CONTACT, "total": len(results), "ok": sum(r["ok"] for r in results), "results": results}


@app.get("/api/registrations", dependencies=[Depends(auth)])
def registrations():
    with db() as c:
        rows = [dict(r) for r in c.execute("SELECT id,created,seller_code,title,price,ok,gd_no,message FROM registrations ORDER BY id DESC LIMIT 100")]
    return rows


@app.get("/api/ai/test", dependencies=[Depends(auth)])
def ai_test():
    """AI 키·모델이 실제로 동작하는지 확인"""
    import requests as rq
    if not AI_KEY:
        return {"ok": False, "message": "ANTHROPIC_API_KEY가 서버에 없어요 (변수 이름 확인 후 재배포)"}
    r = rq.post("https://api.anthropic.com/v1/messages", timeout=30,
                headers={"x-api-key": AI_KEY, "anthropic-version": "2023-06-01", "content-type": "application/json"},
                json={"model": AI_MODEL, "max_tokens": 10, "messages": [{"role": "user", "content": "ping"}]})
    return {"ok": r.status_code == 200, "status": r.status_code, "model": AI_MODEL, "message": r.text[:300]}


@app.get("/api/config", dependencies=[Depends(auth)])
def config():
    return {"demo": DEMO, "ai": bool(AI_KEY), "ai_model": AI_MODEL, "contact": CONTACT}


@app.get("/", dependencies=[Depends(auth)])
def index():
    return FileResponse(Path(__file__).with_name("static") / "index.html")
